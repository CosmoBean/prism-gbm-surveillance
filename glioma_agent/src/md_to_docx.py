"""
Convert Glioma Clinical Report Markdown to DOCX for Google Docs import.
Handles: headings, tables, blockquotes, bold, inline code, images.
"""

import re
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, hex_color: str):
    """Set table cell background color."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)


def add_bold_italic_run(para, text: str):
    """Add a run with inline bold (**text**) and code (`text`) parsing."""
    # Split on **bold** and `code` patterns
    pattern = r'(\*\*[^*]+\*\*|`[^`]+`)'
    parts   = re.split(pattern, text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run      = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run           = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            para.add_run(part)


def parse_pipe_table(lines: list) -> list:
    """Parse Markdown pipe table lines into list of row lists."""
    rows = []
    for line in lines:
        if re.match(r'^\s*\|[-:| ]+\|\s*$', line):
            continue  # skip separator row
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    return rows


def md_to_docx(md_path: str, docx_path: str, images_dir: str):
    doc = Document()

    # ── Page margins ──────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Styles ────────────────────────────────────────────────────────────────
    normal_style          = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)

    with open(md_path, 'r') as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        # ── Heading ───────────────────────────────────────────────────────────
        h_match = re.match(r'^(#{1,4})\s+(.*)', line)
        if h_match:
            level = len(h_match.group(1))
            text  = h_match.group(2).strip()
            # Strip inline formatting from heading
            text  = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
            para  = doc.add_heading(text, level=min(level, 4))
            if level == 1:
                para.runs[0].font.color.rgb = RGBColor(0x1A, 0x5C, 0x8A)
            elif level == 2:
                para.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            i += 1
            continue

        # ── Horizontal rule ───────────────────────────────────────────────────
        if re.match(r'^---+\s*$', line):
            para = doc.add_paragraph()
            pPr  = para._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'),   'single')
            bottom.set(qn('w:sz'),    '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'AAAAAA')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # ── Pipe table ────────────────────────────────────────────────────────
        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[-:| ]+\|\s*$', lines[i + 1]):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].rstrip('\n'))
                i += 1
            rows = parse_pipe_table(table_lines)
            if not rows:
                continue
            n_cols = max(len(r) for r in rows)
            table  = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = 'Table Grid'
            for r_idx, row_data in enumerate(rows):
                for c_idx, cell_text in enumerate(row_data):
                    if c_idx >= n_cols:
                        break
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ''
                    para = cell.paragraphs[0]
                    # Header row styling
                    if r_idx == 0:
                        set_cell_bg(cell, '2E75B6')
                        run = para.add_run(re.sub(r'[\*`]', '', cell_text))
                        run.bold            = True
                        run.font.color.rgb  = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size       = Pt(9.5)
                    else:
                        bg = 'F2F7FC' if r_idx % 2 == 0 else 'FFFFFF'
                        set_cell_bg(cell, bg)
                        add_bold_italic_run(para, cell_text)
                        para.runs[-1].font.size = Pt(9.5) if para.runs else None
            doc.add_paragraph()  # spacing after table
            continue

        # ── Blockquote ────────────────────────────────────────────────────────
        if line.startswith('>'):
            text = line.lstrip('> ').strip()
            para = doc.add_paragraph(style='Normal')
            para.paragraph_format.left_indent  = Inches(0.4)
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after  = Pt(4)
            # Left border via XML
            pPr  = para._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'),   'single')
            left.set(qn('w:sz'),    '12')
            left.set(qn('w:space'), '4')
            left.set(qn('w:color'), '2E75B6')
            pBdr.append(left)
            pPr.append(pBdr)
            add_bold_italic_run(para, text)
            for run in para.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                run.font.italic    = True
            i += 1
            continue

        # ── Image ─────────────────────────────────────────────────────────────
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if img_match:
            img_file = img_match.group(2)
            img_path = os.path.join(images_dir, img_file)
            if os.path.exists(img_path):
                try:
                    doc.add_picture(img_path, width=Inches(6.0))
                    last_para = doc.paragraphs[-1]
                    last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap = doc.add_paragraph(img_match.group(1) or img_file)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.runs[0].font.size   = Pt(9)
                    cap.runs[0].font.italic = True
                    cap.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                except Exception as e:
                    doc.add_paragraph(f"[Image: {img_file}]")
            i += 1
            continue

        # ── Italic note line (*...*) ──────────────────────────────────────────
        if re.match(r'^\*[^*].*[^*]\*\s*$', line):
            text = line.strip().strip('*')
            para = doc.add_paragraph(style='Normal')
            run  = para.add_run(text)
            run.font.italic = True
            run.font.size   = Pt(9)
            run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
            i += 1
            continue

        # ── Empty line ────────────────────────────────────────────────────────
        if not line.strip():
            i += 1
            continue

        # ── Normal paragraph ──────────────────────────────────────────────────
        para = doc.add_paragraph(style='Normal')
        add_bold_italic_run(para, line.strip())
        i += 1

    doc.save(docx_path)
    print(f"✅ DOCX saved: {docx_path}")


if __name__ == "__main__":
    reports_dir = "/home/ubuntu/glioma_agent/reports"

    for tp in ["Timepoint_6", "Timepoint_5"]:
        md_path   = os.path.join(reports_dir, f"PatientID_0019_{tp}_report.md")
        docx_path = os.path.join(reports_dir, f"PatientID_0019_{tp}_report.docx")
        md_to_docx(md_path, docx_path, images_dir=reports_dir)
